"""
routes/vendor_management.py

Endpoints for the Import and Vendor Directory tabs.

POST /api/vendors/import/upload      — import Excel sheet of manual vendors
GET  /api/vendors/master/search      — search vendors_master by name
GET  /api/vendors/master/{id}        — get single vendor
POST /api/vendors/master/add         — add new vendor manually
PATCH /api/vendors/master/{id}       — update/override vendor fields
GET  /api/vendors/master/export/excel — export all master vendors as Excel
GET  /api/vendors/master/stats       — stats for Import tab
POST /api/vendors/master/db-search   — keyword search used by Step 3
"""

import os
import io
import re
import shutil
from datetime import date
from fastapi import APIRouter, HTTPException, UploadFile, File, Form, Depends, Query, Header
from typing import Optional
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional, List

from database import get_suppliers_db

router     = APIRouter(tags=["Vendor Management"])
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)


# ── Models ────────────────────────────────────────────────────

class VendorUpsert(BaseModel):
    vendor_name:             str
    country:                 Optional[str]   = None
    city:                    Optional[str]   = None
    email:                   Optional[str]   = None
    phone:                   Optional[str]   = None
    website:                 Optional[str]   = None
    brands:                  Optional[str]   = None
    description:             Optional[str]   = None
    sourcing_rationale:      Optional[str]   = None
    channel_codes:           Optional[str]   = None
    source:                  Optional[str]   = "manual_override"
    verification_method:     Optional[str]   = "manual"
    verification_confidence: Optional[float] = 1.0
    verification_notes:      Optional[str]   = None


class VendorOverride(BaseModel):
    """Partial update — only fields provided will be updated."""
    vendor_name:          Optional[str]  = None
    country:              Optional[str]  = None
    city:                 Optional[str]  = None
    email:                Optional[str]  = None
    phone:                Optional[str]  = None
    website:              Optional[str]  = None
    brands:               Optional[str]  = None
    description:          Optional[str]  = None
    sourcing_rationale:   Optional[str]  = None
    channel_codes:        Optional[str]  = None
    verification_notes:   Optional[str]  = None
    is_preferred:         Optional[bool] = None
    preferred_divisions:  Optional[str]  = None
    preferred_categories: Optional[str]  = None


# ── Normalize helper ──────────────────────────────────────────

def _normalize(name: str) -> str:
    if not name:
        return ""
    n = name.lower().strip()
    for suffix in [
        r"\bllc\b", r"\bl\.l\.c\b", r"\bltd\b", r"\blimited\b",
        r"\binc\b", r"\bcorp\b", r"\bco\b", r"\bcompany\b",
        r"\bgroup\b", r"\bholding\b", r"\bholdings\b",
        r"\bindustries\b", r"\btrading\b", r"\bfze\b", r"\bfzc\b",
        r"\bsaog\b", r"\bsaoc\b",
    ]:
        n = re.sub(suffix, "", n)
    return re.sub(r"[^\w\s]", "", n).strip()


# ── Upsert helper (used by import + AI pipeline) ──────────────

def upsert_to_master(vendor: dict, db) -> int:
    """
    Insert-or-update a vendor in vendors_master.
    Uses SELECT → UPDATE/INSERT with SAVEPOINT fallback so it works
    regardless of whether the unique index is partial or a full constraint.
    """
    cursor = db.cursor()
    try:
        norm = _normalize(vendor.get("vendor_name", ""))
        if not norm:
            return None

        confidence = float(vendor.get("verification_confidence") or 0.5)
        website    = vendor.get("website") or ""
        if website and not website.startswith("http"):
            website = "https://" + website

        def _do_update(eid, existing_conf):
            if confidence >= existing_conf:
                cursor.execute("""
                    UPDATE vendors_master SET
                        email                   = %s,
                        phone                   = %s,
                        website                 = %s,
                        verification_method     = %s,
                        verification_confidence = GREATEST(%s, verification_confidence),
                        last_verified_date      = GREATEST(
                            %s::date,
                            COALESCE(last_verified_date, '2000-01-01'::date)
                        ),
                        updated_at = NOW()
                    WHERE id = %s
                """, (
                    vendor.get("email"), vendor.get("phone"), website,
                    vendor.get("verification_method", "unknown"),
                    confidence,
                    vendor.get("last_verified_date") or date.today(),
                    eid,
                ))
            cursor.execute("""
                UPDATE vendors_master SET
                    country       = COALESCE(%s, country),
                    city          = COALESCE(%s, city),
                    brands        = COALESCE(%s, brands),
                    description   = COALESCE(%s, description),
                    channel_codes = COALESCE(%s, channel_codes),
                    division_id   = COALESCE(%s, division_id),
                    client        = COALESCE(%s, client),
                    updated_at    = NOW()
                WHERE id = %s
            """, (
                vendor.get("country"), vendor.get("city"),
                vendor.get("brands"), vendor.get("description"),
                vendor.get("channel_codes"),
                vendor.get("division_id"), vendor.get("client"),
                eid,
            ))

        country_val = vendor.get("country", "") or ""

        # 1. Match by (normalized_name, country) — same company different country = separate row
        cursor.execute(
            "SELECT id, verification_confidence FROM vendors_master WHERE normalized_name = %s AND LOWER(COALESCE(country,'')) = LOWER(%s)",
            (norm, country_val)
        )
        existing = cursor.fetchone()

        # 2. Fallback: name+country case-insensitive
        if not existing:
            cursor.execute(
                "SELECT id, verification_confidence FROM vendors_master WHERE LOWER(TRIM(vendor_name)) = LOWER(TRIM(%s)) AND LOWER(COALESCE(country,'')) = LOWER(%s)",
                (vendor.get("vendor_name", ""), country_val)
            )
            existing = cursor.fetchone()

        if existing:
            _do_update(existing[0], float(existing[1] or 0))
            return existing[0]

        # 3. New vendor — INSERT with SAVEPOINT so a race-condition conflict is recoverable
        cursor.execute("SAVEPOINT upsert_vendor")
        try:
            cursor.execute("""
                INSERT INTO vendors_master (
                    vendor_name, normalized_name,
                    country, city, email, phone, website,
                    brands, description, sourcing_rationale, channel_codes,
                    source, discovered_vendor_id,
                    verification_method, verification_confidence,
                    manually_overridden, last_verified_date, verification_notes,
                    division_id, client
                ) VALUES (
                    %s, %s,
                    %s, %s, %s, %s, %s,
                    %s, %s, %s, %s,
                    %s, %s,
                    %s, %s,
                    %s, %s, %s,
                    %s, %s
                ) RETURNING id
            """, (
                vendor.get("vendor_name"), norm,
                vendor.get("country"), vendor.get("city"),
                vendor.get("email"), vendor.get("phone"), website,
                vendor.get("brands"), vendor.get("description"),
                vendor.get("sourcing_rationale"), vendor.get("channel_codes"),
                vendor.get("source", "unknown"), vendor.get("discovered_vendor_id"),
                vendor.get("verification_method", "unknown"), confidence,
                vendor.get("manually_overridden", False),
                vendor.get("last_verified_date") or date.today(),
                vendor.get("verification_notes"),
                vendor.get("division_id"), vendor.get("client"),
            ))
            row = cursor.fetchone()
            cursor.execute("RELEASE SAVEPOINT upsert_vendor")
            return row[0] if row else None
        except Exception:
            # Unique constraint fired — roll back to savepoint and update instead
            cursor.execute("ROLLBACK TO SAVEPOINT upsert_vendor")
            cursor.execute(
                "SELECT id, verification_confidence FROM vendors_master WHERE normalized_name = %s OR LOWER(TRIM(vendor_name)) = LOWER(TRIM(%s))",
                (norm, vendor.get("vendor_name", ""))
            )
            existing2 = cursor.fetchone()
            if existing2:
                _do_update(existing2[0], float(existing2[1] or 0))
                return existing2[0]
            raise

    finally:
        cursor.close()


# ── Feature 1: Import Excel ───────────────────────────────────

COLUMN_ALIASES = {
    "vendor_name": [
        "vendor name", "company", "company name", "supplier", "supplier name",
        "name", "vendor",
    ],
    "country": [
        "country", "country/region", "location", "region",
    ],
    "city": [
        "city", "town", "emirate",
    ],
    "email": [
        "email", "email address", "e-mail", "contact email", "contact e-mail",
        "official email", "official email address",
        "business email", "working official business email",
    ],
    "phone": [
        "phone", "telephone", "mobile", "contact phone", "tel", "phone number",
        "official phone", "official phone number",
    ],
    "website": [
        "website", "web", "url", "web site", "web address",
    ],
    "brands": [
        "brands", "products", "product brands", "brands handled",
        "products / brands handled", "products/brands", "brands / products",
        "brands/products",
    ],
    "description": [
        "description", "notes", "remarks", "about",
    ],
    "channel_codes": [
        "channel", "channels", "channel code", "supplier type",
        "supplier type (group a channel)",
    ],
    "rfq_probability": [
        "rfq response probability", "response probability", "probability",
        "rfq probability",
    ],
    "speed_rating": [
        "quotation speed rating", "speed rating", "quotation speed",
        "response speed", "rating",
    ],
}

_EMPTY_MARKERS = {
    "nan", "none", "", "n/a", "na", "nil", "-",
    "contact via website", "contact via distributor locator",
    "contact via website.", "see website", "check website",
    "verify on site", "via website",
}

_SEARCH_MARKERS = ("search ", "find ", "look up ")

SPEED_TO_CONFIDENCE = {
    "very high": 1.00,
    "high":      0.85,
    "medium":    0.65,
    "low":       0.45,
}


def _probability_to_confidence(val: str) -> Optional[float]:
    try:
        p = float(val)
        return round(p / 100, 2) if 0 <= p <= 100 else None
    except (ValueError, TypeError):
        return None


def _clean_field(val: str, field: str = "") -> Optional[str]:
    if not val:
        return None
    v = val.strip()
    if v.lower() in _EMPTY_MARKERS:
        return None
    if field == "phone":
        if any(v.lower().startswith(m) for m in _SEARCH_MARKERS):
            return None
        if not any(c.isdigit() for c in v):
            return None
    if field == "website":
        if any(v.lower().startswith(m) for m in _SEARCH_MARKERS):
            return None
    return v if v else None


def _clean_emails(raw: str) -> str:
    if not raw:
        return ""
    parts         = re.split(r'\s*/\s*|\s*;\s*|\s*,\s*', raw.strip())
    email_pattern = re.compile(r'^[^\s@]+@[^\s@]+\.[^\s@]+$')
    valid         = []
    for p in parts:
        p = p.strip()
        if p and email_pattern.match(p) and p.lower() not in _EMPTY_MARKERS:
            valid.append(p.lower())
    return " / ".join(valid)


def _detect_column(headers: List[str], field: str) -> Optional[int]:
    aliases = COLUMN_ALIASES.get(field, [field])
    for i, h in enumerate(headers):
        h_clean = h.strip().lower()
        for alias in aliases:
            if alias in h_clean or h_clean in alias:
                return i
    return None


def _infer_product_from_filename(filename: str) -> str:
    name = os.path.splitext(filename)[0]
    name = re.sub(r'^[A-Z0-9]+[-_]+[-_]+', '', name)
    name = re.sub(r'[-_]+', ' ', name).strip()
    return name


@router.post("/vendors/import/upload")
async def import_vendors_excel(
    file:                UploadFile = File(...),
    product_description: str        = Form(""),
    division_id:         str        = Form(""),
    client:              str        = Form(""),
    db=Depends(get_suppliers_db),
    x_username:  Optional[str] = Header(None),
    x_user_role: Optional[str] = Header(None),
):
    """
    Upload an Excel sheet of manually verified vendors.
    Manual imports go directly to vendors_master only —
    no discovered_vendors entry needed since these are
    standalone clean lists, not inquiry-linked searches.
    """
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in {".xlsx", ".xls", ".xlsm", ".csv"}:
        raise HTTPException(status_code=400,
            detail="Only Excel (.xlsx, .xls, .xlsm) and CSV files supported")

    product_context = (
        product_description.strip()
        or _infer_product_from_filename(file.filename or "")
    )

    file_path = os.path.join(UPLOAD_DIR, file.filename)
    try:
        with open(file_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        try:
            import pandas as pd
            if ext == ".csv":
                df = pd.read_csv(file_path, dtype=str)
            else:
                df = pd.read_excel(file_path, dtype=str)
            df = df.fillna("").head(5000)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not read file: {e}")

        raw_headers = [str(c).strip() for c in df.columns]
        col_map     = {}
        for field in COLUMN_ALIASES.keys():
            idx = _detect_column(raw_headers, field)
            if idx is not None:
                col_map[field] = idx

        if "vendor_name" not in col_map:
            raise HTTPException(status_code=400,
                detail=f"Could not find a 'Vendor Name' / 'Company Name' column. "
                       f"Columns found: {raw_headers}")

        saved       = 0
        updated     = 0
        skipped     = 0
        errors      = []
        sample_rows = []

        # Pre-create the import log entry so we get an ID to stamp on each vendor
        import_log_id = None
        try:
            pre_cur = db.cursor()
            pre_cur.execute(
                "INSERT INTO vendor_import_log (filename, saved, updated, skipped, total_rows) VALUES (%s, 0, 0, 0, %s) RETURNING id",
                (file.filename, len(df))
            )
            import_log_id = pre_cur.fetchone()[0]
            db.commit()
            pre_cur.close()
        except Exception:
            pass

        for row_idx, row in df.iterrows():
            row_vals = [str(v).strip() for v in row.values]

            def get_raw(f, _rv=row_vals, _cm=col_map):
                idx = _cm.get(f)
                if idx is None or idx >= len(_rv):
                    return None
                v = _rv[idx]
                return v if v and v.lower() not in ("nan", "none", "") else None

            name = _clean_field(get_raw("vendor_name"))
            if not name:
                skipped += 1
                continue

            email_cleaned = _clean_emails(get_raw("email") or "")

            confidence = 1.0
            raw_prob   = get_raw("rfq_probability")
            raw_speed  = get_raw("speed_rating")
            if raw_prob:
                p = _probability_to_confidence(raw_prob)
                if p is not None:
                    confidence = p
            elif raw_speed:
                confidence = SPEED_TO_CONFIDENCE.get(
                    raw_speed.lower().strip(), 1.0
                )

            notes_parts = []
            if raw_speed:
                notes_parts.append(f"Quotation Speed: {raw_speed}")
            if raw_prob:
                notes_parts.append(f"RFQ Response Probability: {raw_prob}%")
            notes_parts.append("Manually imported via Excel upload")

            vendor = {
                "vendor_name":             name,
                "country":                 _clean_field(get_raw("country")),
                "city":                    _clean_field(get_raw("city")),
                "email":                   email_cleaned or None,
                "phone":                   _clean_field(get_raw("phone"), "phone"),
                "website":                 _clean_field(get_raw("website"), "website"),
                "brands":                  _clean_field(get_raw("brands")),
                "description":             _clean_field(get_raw("description")),
                "sourcing_rationale":      product_context,
                "channel_codes":           _clean_field(get_raw("channel_codes")),
                "source":                  "manual_import",
                "verification_method":     "manual",
                "verification_confidence": confidence,
                "manually_overridden":     False,
                "discovered_vendor_id":    None,
                "last_verified_date":      date.today(),
                "verification_notes":      " | ".join(notes_parts),
                "division_id":             int(division_id) if division_id.strip().isdigit() else None,
                "client":                  client.strip() or None,
            }

            try:
                norm      = _normalize(name)
                cur       = db.cursor()

                # Match by (name, country) — same company in different country gets its own row
                v_country = vendor.get("country", "") or ""
                cur.execute(
                    "SELECT id FROM vendors_master WHERE (normalized_name = %s OR LOWER(TRIM(vendor_name)) = LOWER(TRIM(%s))) AND LOWER(COALESCE(country,'')) = LOWER(%s) LIMIT 1",
                    (norm, name, v_country)
                )
                existing_row = cur.fetchone()

                if existing_row:
                    # UPDATE existing vendor
                    eid = existing_row[0]
                    cur.execute("""
                        UPDATE vendors_master SET
                            email         = COALESCE(%s, email),
                            phone         = COALESCE(%s, phone),
                            website       = COALESCE(%s, website),
                            country       = COALESCE(%s, country),
                            city          = COALESCE(%s, city),
                            brands        = COALESCE(%s, brands),
                            division_id   = COALESCE(%s, division_id),
                            client        = COALESCE(%s, client),
                            import_log_id = COALESCE(%s, import_log_id),
                            updated_at    = NOW()
                        WHERE id = %s
                    """, (
                        vendor["email"], vendor["phone"], vendor["website"],
                        vendor["country"], vendor["city"], vendor["brands"],
                        vendor["division_id"], vendor["client"],
                        import_log_id, eid,
                    ))
                    db.commit()
                    cur.close()
                    updated += 1
                else:
                    # INSERT new vendor — use SAVEPOINT so conflict is recoverable
                    cur.execute("SAVEPOINT imp_sp")
                    try:
                        cur.execute("""
                            INSERT INTO vendors_master (
                                vendor_name, normalized_name,
                                country, city, email, phone, website,
                                brands, description, sourcing_rationale, channel_codes,
                                source, verification_method, verification_confidence,
                                manually_overridden, last_verified_date, verification_notes,
                                division_id, client, import_log_id
                            ) VALUES (
                                %s, %s,
                                %s, %s, %s, %s, %s,
                                %s, %s, %s, %s,
                                %s, %s, %s,
                                %s, %s, %s,
                                %s, %s, %s
                            )
                        """, (
                            vendor["vendor_name"], norm,
                            vendor["country"], vendor["city"],
                            vendor["email"], vendor["phone"], vendor["website"],
                            vendor["brands"], vendor["description"],
                            vendor["sourcing_rationale"], vendor["channel_codes"],
                            vendor["source"], vendor["verification_method"],
                            vendor["verification_confidence"],
                            vendor["manually_overridden"],
                            vendor["last_verified_date"] or date.today(),
                            vendor["verification_notes"],
                            vendor["division_id"], vendor["client"], import_log_id,
                        ))
                        cur.execute("RELEASE SAVEPOINT imp_sp")
                        db.commit()
                        cur.close()
                        saved += 1
                        if len(sample_rows) < 5:
                            sample_rows.append({
                                "name":    vendor["vendor_name"],
                                "country": vendor["country"],
                                "email":   vendor["email"],
                                "phone":   vendor["phone"],
                            })
                    except Exception as insert_err:
                        cur.execute("ROLLBACK TO SAVEPOINT imp_sp")
                        # Conflict on insert — try update by name
                        cur.execute(
                            "SELECT id FROM vendors_master WHERE normalized_name = %s OR LOWER(TRIM(vendor_name)) = LOWER(TRIM(%s)) LIMIT 1",
                            (norm, name)
                        )
                        fallback = cur.fetchone()
                        if fallback:
                            cur.execute("""
                                UPDATE vendors_master SET
                                    email         = COALESCE(%s, email),
                                    phone         = COALESCE(%s, phone),
                                    website       = COALESCE(%s, website),
                                    country       = COALESCE(%s, country),
                                    city          = COALESCE(%s, city),
                                    brands        = COALESCE(%s, brands),
                                    division_id   = COALESCE(%s, division_id),
                                    client        = COALESCE(%s, client),
                                    import_log_id = COALESCE(%s, import_log_id),
                                    updated_at    = NOW()
                                WHERE id = %s
                            """, (
                                vendor["email"], vendor["phone"], vendor["website"],
                                vendor["country"], vendor["city"], vendor["brands"],
                                vendor["division_id"], vendor["client"],
                                import_log_id, fallback[0],
                            ))
                            db.commit()
                            cur.close()
                            updated += 1
                        else:
                            cur.close()
                            db.rollback()
                            errors.append(f"Row {row_idx+1} ({name}): {str(insert_err)[:120]}")
                        continue

            except Exception as e:
                db.rollback()
                errors.append(f"Row {row_idx+1} ({name}): {str(e)[:120]}")
                continue

        # Update the pre-created log entry with final counts
        try:
            log_cur = db.cursor()
            if import_log_id:
                log_cur.execute(
                    "UPDATE vendor_import_log SET saved=%s, updated=%s, skipped=%s, total_rows=%s WHERE id=%s",
                    (saved, updated, skipped, len(df), import_log_id)
                )
            else:
                log_cur.execute(
                    "INSERT INTO vendor_import_log (filename, saved, updated, skipped, total_rows) VALUES (%s,%s,%s,%s,%s)",
                    (file.filename, saved, updated, skipped, len(df))
                )
            db.commit()
            log_cur.close()
        except Exception:
            pass

        # Log to activity feed
        try:
            from routes.history import log_activity
            log_activity(
                event_type="vendor_import",
                title=f"Vendor Import — {file.filename}",
                detail=f"{saved} new, {updated} updated, {skipped} skipped out of {len(df)} rows",
                meta={"filename": file.filename, "saved": saved, "updated": updated, "skipped": skipped, "total_rows": len(df)},
                username=x_username,
                user_role=x_user_role,
            )
        except Exception:
            pass

        return {
            "success":       True,
            "filename":      file.filename,
            "total_rows":    len(df),
            "saved":         saved,
            "updated":       updated,
            "skipped":       skipped,
            "errors":        errors,
            "columns_found": list(col_map.keys()),
            "sample":        sample_rows,
            "message":       f"✅ {saved} new + {updated} updated vendors processed successfully.",
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)


# ── Search vendors_master ─────────────────────────────────────

@router.get("/vendors/master/search")
def search_vendors_master(
    q:              str            = Query("", description="Search by name, email, or brand"),
    source:         Optional[str]  = Query(None),
    country:        Optional[str]  = Query(None),
    date_from:      Optional[str]  = Query(None),
    date_to:        Optional[str]  = Query(None),
    sort:           str            = Query("date_desc"),
    limit:          int            = Query(50, le=200),
    offset:         int            = Query(0, ge=0),
    preferred_only: bool           = Query(False),
    db=Depends(get_suppliers_db),
):
    cursor = db.cursor()
    try:
        conditions = []
        params     = []

        if q.strip():
            conditions.append("""(
                LOWER(vendor_name)             ILIKE %s OR
                LOWER(COALESCE(email,''))       ILIKE %s OR
                LOWER(COALESCE(brands,''))      ILIKE %s OR
                LOWER(COALESCE(city,''))        ILIKE %s
            )""")
            like = f"%{q.lower()}%"
            params.extend([like, like, like, like])

        if source:
            conditions.append("source = %s")
            params.append(source)

        if country:
            conditions.append("LOWER(country) ILIKE %s")
            params.append(f"%{country.lower()}%")

        if date_from:
            conditions.append("created_at::date >= %s")
            params.append(date_from)

        if date_to:
            conditions.append("created_at::date <= %s")
            params.append(date_to)

        if preferred_only:
            conditions.append("is_preferred = TRUE")

        where = ("WHERE " + " AND ".join(conditions)) if conditions else ""

        order_by = {
            "date_desc": "created_at DESC",
            "date_asc":  "created_at ASC",
            "name_asc":  "vendor_name ASC",
            "name_desc": "vendor_name DESC",
        }.get(sort, "created_at DESC")

        cursor.execute(f"""
            SELECT id, vendor_name, country, city, email, phone, website,
                   brands, description, channel_codes, source,
                   verification_method, verification_confidence,
                   manually_overridden, last_verified_date::TEXT,
                   verification_notes, created_at::TEXT, updated_at::TEXT,
                   is_preferred, preferred_divisions, preferred_categories
            FROM vendors_master
            {where}
            ORDER BY is_preferred DESC NULLS LAST, {order_by}
            LIMIT %s OFFSET %s
        """, params + [limit, offset])

        cols    = [d[0] for d in cursor.description]
        results = [dict(zip(cols, row)) for row in cursor.fetchall()]

        cursor.execute(f"SELECT COUNT(*) FROM vendors_master {where}", params)
        total = cursor.fetchone()[0]

        return {"success": True, "vendors": results, "total": total, "showing": len(results)}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.get("/vendors/master/stats")
def get_master_stats(db=Depends(get_suppliers_db)):
    cursor = db.cursor()
    try:
        cursor.execute("""
            SELECT
                COUNT(*)                                                    AS total,
                COUNT(*) FILTER (WHERE source = 'manual_import')           AS manual_imports,
                COUNT(*) FILTER (WHERE source = 'ai_discovered')           AS ai_discovered,
                COUNT(*) FILTER (WHERE source = 'manual_override')         AS manual_overrides,
                COUNT(*) FILTER (WHERE verification_confidence >= 0.9)     AS high_confidence,
                COUNT(*) FILTER (WHERE manually_overridden = TRUE)         AS overridden,
                AVG(verification_confidence)::NUMERIC(4,2)                 AS avg_confidence
            FROM vendors_master
        """)
        cols  = [d[0] for d in cursor.description]
        row   = cursor.fetchone()
        stats = dict(zip(cols, row))
        for k, v in stats.items():
            if hasattr(v, '__float__'):
                try:
                    stats[k] = float(v)
                except Exception:
                    pass
        return {"success": True, "stats": stats}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.get("/vendors/master/{vendor_id}")
def get_vendor(vendor_id: int, db=Depends(get_suppliers_db)):
    cursor = db.cursor()
    try:
        cursor.execute("""
            SELECT id, vendor_name, country, city, email, phone, website,
                   brands, description, sourcing_rationale, channel_codes,
                   source, verification_method, verification_confidence,
                   manually_overridden, last_verified_date::TEXT,
                   verification_notes, created_at::TEXT, updated_at::TEXT
            FROM vendors_master WHERE id = %s
        """, (vendor_id,))
        cols = [d[0] for d in cursor.description]
        row  = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Vendor not found")
        return {"success": True, "vendor": dict(zip(cols, row))}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.post("/vendors/master/add")
def add_vendor_manually(req: VendorUpsert, db=Depends(get_suppliers_db)):
    if not req.vendor_name.strip():
        raise HTTPException(status_code=400, detail="vendor_name is required")
    try:
        vendor = {
            **req.dict(),
            "source":                  "manual_override",
            "verification_method":     "manual",
            "verification_confidence": 1.0,
            "manually_overridden":     True,
            "discovered_vendor_id":    None,
            "last_verified_date":      date.today(),
            "verification_notes":      "Manually added via Vendor Directory",
        }
        master_id = upsert_to_master(vendor, db)
        db.commit()
        return {"success": True, "id": master_id, "message": "Vendor added successfully"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.patch("/vendors/master/{vendor_id}")
def override_vendor(
    vendor_id: int,
    req:       VendorOverride,
    db=Depends(get_suppliers_db),
):
    cursor = db.cursor()
    try:
        updates = {
            field: val
            for field, val in req.dict(exclude_none=True).items()
            if val is not None and val != ""
        }

        if not updates:
            raise HTTPException(status_code=400, detail="No fields provided to update")

        updates["manually_overridden"]     = True
        updates["verification_confidence"] = 1.0
        updates["verification_method"]     = "manual"
        updates["last_verified_date"]      = date.today()

        set_parts = []
        params    = []
        for k, v in updates.items():
            set_parts.append(f"{k} = %s")
            params.append(v)
        set_parts.append("updated_at = NOW()")

        params.append(vendor_id)
        cursor.execute(
            f"UPDATE vendors_master SET {', '.join(set_parts)} WHERE id = %s RETURNING id",
            params,
        )
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Vendor not found")

        db.commit()
        return {"success": True, "message": "Vendor updated successfully"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.post("/vendors/master/{vendor_id}/toggle-preferred")
def toggle_preferred(vendor_id: int, req: dict = None, db=Depends(get_suppliers_db)):
    """Toggle is_preferred flag and optionally set divisions/categories scope."""
    if req is None:
        req = {}
    cursor = db.cursor()
    try:
        cursor.execute("SELECT is_preferred FROM vendors_master WHERE id = %s", (vendor_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Vendor not found")
        new_val = not bool(row[0])
        cursor.execute("""
            UPDATE vendors_master SET
                is_preferred         = %s,
                preferred_divisions  = %s,
                preferred_categories = %s,
                preferred_marked_at  = CASE WHEN %s THEN NOW() ELSE NULL END,
                updated_at           = NOW()
            WHERE id = %s
        """, (
            new_val,
            req.get("preferred_divisions") if new_val else None,
            req.get("preferred_categories") if new_val else None,
            new_val,
            vendor_id,
        ))
        db.commit()
        return {"success": True, "is_preferred": new_val}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.post("/vendors/master/bulk-delete")
def bulk_delete_vendors(payload: dict, db=Depends(get_suppliers_db)):
    """Delete multiple vendors in one query."""
    ids = [int(i) for i in payload.get("ids", []) if str(i).isdigit()]
    if not ids:
        raise HTTPException(status_code=400, detail="No valid ids provided")
    cursor = db.cursor()
    try:
        ph = ",".join(["%s"] * len(ids))
        cursor.execute(f"DELETE FROM vendors_master WHERE id IN ({ph}) RETURNING id", ids)
        deleted = len(cursor.fetchall())
        db.commit()
        return {"success": True, "deleted": deleted}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.delete("/vendors/import-log/{log_id}")
def delete_import_log(log_id: int, db=Depends(get_suppliers_db)):
    cursor = db.cursor()
    try:
        # Delete vendors that belong to this import batch
        cursor.execute(
            "DELETE FROM vendors_master WHERE import_log_id = %s",
            (log_id,)
        )
        vendors_deleted = cursor.rowcount

        # Delete the log entry itself
        cursor.execute("DELETE FROM vendor_import_log WHERE id = %s RETURNING id", (log_id,))
        row = cursor.fetchone()
        db.commit()
        if not row:
            raise HTTPException(status_code=404, detail="Import log entry not found")
        return {"success": True, "vendors_deleted": vendors_deleted}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.delete("/vendors/master/{vendor_id}")
def delete_vendor(vendor_id: int, db=Depends(get_suppliers_db)):
    """Permanently delete a vendor from vendors_master."""
    cursor = db.cursor()
    try:
        cursor.execute("DELETE FROM vendors_master WHERE id = %s RETURNING id", (vendor_id,))
        if not cursor.fetchone():
            raise HTTPException(status_code=404, detail="Vendor not found")
        db.commit()
        return {"success": True, "message": "Vendor deleted"}
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.delete("/vendors/delete-by-inquiry/{inquiry_id}")
def delete_vendors_by_inquiry(inquiry_id: str, db=Depends(get_suppliers_db)):
    """Delete all vendor search results and the inquiry record for a given inquiry_id."""
    cursor = db.cursor()
    try:
        # 1. Delete discovered_vendors linked directly to this inquiry_id
        cursor.execute("DELETE FROM discovered_vendors WHERE inquiry_id = %s", (inquiry_id,))
        deleted_vendors = cursor.rowcount

        # 2. Delete search log entries for this inquiry
        cursor.execute("DELETE FROM inquiry_search_log WHERE inquiry_id = %s", (inquiry_id,))

        # 3. Delete line items if multi-item inquiry
        cursor.execute(
            "DELETE FROM inquiry_line_items WHERE inquiry_id = %s",
            (inquiry_id,)
        )

        # 4. Delete from inquiry_master (multi-item) — ignore if not found
        cursor.execute("DELETE FROM inquiry_master WHERE inquiry_id = %s", (inquiry_id,))

        # 5. Delete from inquiry_history (single-item) — ignore if not found
        cursor.execute("DELETE FROM inquiry_history WHERE inquiry_id = %s", (inquiry_id,))

        db.commit()
        return {
            "success": True,
            "deleted_vendors": deleted_vendors,
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.get("/vendors/master/export/excel")
def export_vendors_master(
    source: Optional[str] = Query(None),
    db=Depends(get_suppliers_db),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    cursor = db.cursor()
    try:
        where  = "WHERE source = %s" if source else ""
        params = [source] if source else []
        cursor.execute(f"""
            SELECT vendor_name, country, city, email, phone, website,
                   brands, channel_codes, source, verification_method,
                   verification_confidence, manually_overridden,
                   last_verified_date::TEXT, verification_notes,
                   created_at::TEXT
            FROM vendors_master
            {where}
            ORDER BY verification_confidence DESC, vendor_name
        """, params)
        rows = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vendors Master"

        headers = [
            "Vendor Name", "Country", "City", "Email", "Phone", "Website",
            "Brands", "Channels", "Source", "Verification Method",
            "Confidence", "Manually Overridden", "Last Verified", "Notes", "Created At",
        ]
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor="1F3864")
        for ci, h in enumerate(headers, 1):
            cell           = ws.cell(row=1, column=ci, value=h)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 20

        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        label    = f"_{source}" if source else ""
        filename = f"vendors_master{label}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


# ── Step 3 DB search endpoint ─────────────────────────────────

@router.post("/vendors/master/db-search")
def search_vendors_master_for_inquiry(
    req: dict,
    db=Depends(get_suppliers_db),
):
    """
    Searches vendors_master by keywords — used in Step 3 DB search
    alongside Suppliers_Clean_Phase1.
    sourcing_rationale is the key field: it stores what product the
    vendor was originally found for (set at import time).
    """
    keywords      = str(req.get("keywords", "")).strip()
    brand         = str(req.get("brand",    "") or "").strip()
    top_n         = int(req.get("top_n", 15))

    if not keywords:
        return {"success": True, "results": [], "total_found": 0}

    cursor = db.cursor()
    try:
        words = [w for w in keywords.split() if len(w) >= 3][:6]
        if not words:
            words = keywords.split()[:3]

        conditions = []
        params     = []
        for word in words:
            w = f"%{word.lower()}%"
            conditions.append("""(
                LOWER(vendor_name)                      ILIKE %s OR
                LOWER(COALESCE(brands,''))               ILIKE %s OR
                LOWER(COALESCE(sourcing_rationale,''))   ILIKE %s OR
                LOWER(COALESCE(description,''))          ILIKE %s OR
                LOWER(COALESCE(channel_codes,''))        ILIKE %s
            )""")
            params.extend([w, w, w, w, w])

        where = "WHERE (" + " OR ".join(conditions) + ")"
        if brand:
            where += " AND LOWER(COALESCE(brands,'')) ILIKE %s"
            params.append(f"%{brand.lower()}%")

        cursor.execute(f"""
            SELECT vendor_name, country, city, email, phone, website,
                   brands, sourcing_rationale, description, channel_codes,
                   verification_confidence, verification_method,
                   last_verified_date::TEXT, source
            FROM vendors_master
            {where}
            ORDER BY verification_confidence DESC, updated_at DESC
            LIMIT %s
        """, params + [top_n * 2])

        cols = [d[0] for d in cursor.description]
        rows = [dict(zip(cols, row)) for row in cursor.fetchall()]

        desc_words = set(w.lower() for w in keywords.split() if len(w) >= 3)
        scored     = []
        for row in rows:
            text = " ".join([
                row.get("vendor_name",        "") or "",
                row.get("brands",             "") or "",
                row.get("sourcing_rationale", "") or "",
                row.get("description",        "") or "",
            ]).lower()
            kw_hits = sum(1 for w in desc_words if w in text)
            conf    = float(row.get("verification_confidence") or 0.5)
            score   = kw_hits * 3 + conf * 5
            if score > 0:
                scored.append({
                    **row,
                    "score":         round(score, 1),
                    "match_reasons": [f"vendors_master ({kw_hits} keyword hits)"],
                })

        scored.sort(key=lambda x: x["score"], reverse=True)
        results = scored[:top_n]

        return {
            "success":       True,
            "results":       results,
            "total_found":   len(results),
            "query_summary": f"Found {len(results)} vendors in master DB for: {keywords}",
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


# ── Update product description for imported batches ───────────

@router.get("/vendors/master/product-contexts")
def get_product_contexts(db=Depends(get_suppliers_db)):
    """
    Returns all distinct sourcing_rationale values with vendor counts.
    Used in the Import tab to show existing product descriptions
    that can be corrected.
    """
    cursor = db.cursor()
    try:
        cursor.execute("""
            SELECT
                sourcing_rationale,
                COUNT(*) AS vendor_count,
                MAX(created_at)::TEXT AS last_imported
            FROM vendors_master
            WHERE sourcing_rationale IS NOT NULL
              AND sourcing_rationale != ''
              AND source = 'manual_import'
            GROUP BY sourcing_rationale
            ORDER BY MAX(created_at) DESC
        """)
        cols = [d[0] for d in cursor.description]
        rows = [dict(zip(cols, row)) for row in cursor.fetchall()]
        return {"success": True, "contexts": rows}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


@router.patch("/vendors/master/update-product-context")
def update_product_context(
    req: dict,
    db=Depends(get_suppliers_db),
):
    """
    Updates sourcing_rationale for all vendors matching old_description.
    Used when the wrong product description was used during import.
    """
    old_description = (req.get("old_description") or "").strip()
    new_description = (req.get("new_description") or "").strip()

    if not old_description:
        raise HTTPException(status_code=400, detail="old_description is required")
    if not new_description:
        raise HTTPException(status_code=400, detail="new_description is required")
    if old_description == new_description:
        raise HTTPException(status_code=400, detail="New description is the same as old")

    cursor = db.cursor()
    try:
        cursor.execute("""
            UPDATE vendors_master
            SET sourcing_rationale = %s,
                updated_at         = NOW()
            WHERE sourcing_rationale = %s
            RETURNING id
        """, (new_description, old_description))

        updated = len(cursor.fetchall())
        db.commit()

        return {
            "success": True,
            "updated": updated,
            "message": f"✅ Updated product description for {updated} vendors.",
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


# ── Retag: get distinct sourcing_rationale values ─────────────

@router.get("/vendors/master/import-descriptions")
def get_import_descriptions(db=Depends(get_suppliers_db)):
    """
    Returns distinct sourcing_rationale values from manual_import vendors
    with a count of how many vendors have each description.
    Used by the retag dropdown in the Import tab.
    """
    cursor = db.cursor()
    try:
        cursor.execute("""
            SELECT
                sourcing_rationale  AS description,
                COUNT(*)            AS count
            FROM vendors_master
            WHERE source = 'manual_import'
              AND sourcing_rationale IS NOT NULL
              AND sourcing_rationale != ''
            GROUP BY sourcing_rationale
            ORDER BY COUNT(*) DESC, sourcing_rationale
        """)
        rows = cursor.fetchall()
        return {
            "success":      True,
            "descriptions": [{"description": r[0], "count": r[1]} for r in rows],
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()


# ── Retag: bulk update sourcing_rationale ─────────────────────

class RetagRequest(BaseModel):
    old_description: str
    new_description: str


@router.post("/vendors/master/sync-from-discovered")
def sync_discovered_to_master(db=Depends(get_suppliers_db)):
    """
    Copies any AI-discovered vendors that are not yet in vendors_master.
    Runs upsert_to_master for each missing one. Safe to call multiple times.
    """
    cursor = db.cursor()
    try:
        cursor.execute("""
            SELECT DISTINCT ON (LOWER(dv.vendor_name))
                dv.vendor_name, dv.country, dv.city, dv.email, dv.phone,
                dv.website, dv.brands, dv.description, dv.source,
                im.division_id, im.client,
                STRING_AGG(sc.code, ', ' ORDER BY sc.code) OVER (PARTITION BY dv.id) AS channel_codes
            FROM discovered_vendors dv
            LEFT JOIN inquiry_master im ON im.inquiry_id = dv.inquiry_id
            LEFT JOIN vendor_channel_map vcm ON vcm.vendor_id = dv.id
            LEFT JOIN sourcing_channels sc ON sc.id = vcm.channel_id
            WHERE LOWER(dv.vendor_name) NOT IN (
                SELECT LOWER(vendor_name) FROM vendors_master WHERE vendor_name IS NOT NULL
            )
            ORDER BY LOWER(dv.vendor_name), dv.id DESC
        """)
        rows = cursor.fetchall()
        cursor.close()

        synced = 0
        failed = 0
        for r in rows:
            try:
                upsert_to_master({
                    "vendor_name":             r[0],
                    "country":                 r[1],
                    "city":                    r[2],
                    "email":                   r[3],
                    "phone":                   r[4],
                    "website":                 r[5] or "",
                    "brands":                  r[6],
                    "description":             r[7],
                    "sourcing_rationale":      r[7],
                    "channel_codes":           r[11],
                    "source":                  "ai_discovered",
                    "verification_method":     r[8] or "ai_pipeline",
                    "verification_confidence": 0.7,
                    "manually_overridden":     False,
                    "last_verified_date":      None,
                    "verification_notes":      f"Synced from discovered_vendors",
                    "division_id":             r[9],
                    "client":                  r[10],
                }, db)
                synced += 1
            except Exception as e:
                print(f"  ⚠️  sync failed for {r[0]}: {e}")
                failed += 1

        return {"success": True, "synced": synced, "failed": failed}
    except Exception as e:
        cursor.close()
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/vendors/master/retag-description")
def retag_description(req: RetagRequest, db=Depends(get_suppliers_db)):
    """
    Bulk updates sourcing_rationale for all vendors where it matches
    old_description. Used to fix wrong product descriptions on imports.
    """
    if not req.old_description.strip():
        raise HTTPException(status_code=400, detail="old_description is required")
    if not req.new_description.strip():
        raise HTTPException(status_code=400, detail="new_description is required")
    if req.old_description.strip() == req.new_description.strip():
        raise HTTPException(status_code=400, detail="Old and new descriptions are the same")

    cursor = db.cursor()
    try:
        cursor.execute("""
            UPDATE vendors_master
            SET sourcing_rationale = %s,
                updated_at         = NOW()
            WHERE sourcing_rationale = %s
              AND source = 'manual_import'
        """, (req.new_description.strip(), req.old_description.strip()))

        updated = cursor.rowcount
        db.commit()

        return {
            "success": True,
            "updated": updated,
            "message": f"Updated {updated} vendors from '{req.old_description}' → '{req.new_description}'",
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        cursor.close()